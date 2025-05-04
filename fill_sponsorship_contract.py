from docx import Document

def fill_contract(data, template_path="sponsorship_agreement_template.docx", output_path="filled_sponsorship_agreement.docx"):
    doc = Document(template_path)

    placeholders = {
        "[Insert Date]": data.get("Date", ""),
        "[Sponsor Name]": data.get("Sponsor", ""),
        "[Email]": data.get("Email", ""),
        "[Insert Address]": data.get("Address", ""),
        "[Facility or Organization Name]": data.get("Facility", ""),
        "[Facility Address]": data.get("Facility Address", ""),
        "[Asset Type]": data.get("Asset Type", ""),
        "[Location Scope]": data.get("Location", ""),
        "[Tier Level]": data.get("Tier", ""),
        "[Yes/No]": "Yes" if data.get("Exclusivity", False) else "No",
        "[Impressions]": str(data.get("Impressions", "")),
        "[Duration Months]": str(data.get("Duration", "")),
        "[Final Price]": f"{data.get('Price', 0):,.2f}"
    }

    for para in doc.paragraphs:
        for placeholder, value in placeholders.items():
            if placeholder in para.text:
                para.text = para.text.replace(placeholder, str(value))

    doc.save(output_path)
    return output_path
